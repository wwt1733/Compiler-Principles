# -*- coding: utf-8 -*-
"""课程设计三份文档（实验报告书 / 系统使用说明书 / 测试报告）的公共排版库。

排版依据：
  * reference/taskbook/编译原理课程设计实验报告书 版头及正文格式.doc
      A4（595.3 x 841.9 pt），左右页边距 90 pt（≈3.17 cm），上下 72 pt（≈2.54 cm）；
      正文：宋体 小四（12 pt），固定行距 18 pt，两端对齐；
      章节：一、…／1.1／1.1.1 三级标题。
  * reference/taskbook/编译原理课程项目 实验报告书 书写格式 模板 范例.pdf
      每个功能按“存储结构 / 算法实现过程文字描述 / 算法实现流程图 /
      关键源程序代码段 / 单元测试”五段书写；图、表均带“图 1-1 / 表 1-1”式编号与标题。
  * reference/taskbook/算法流程图的画法-含错误示例.pdf
      流程图单入口单出口、只用顺序/判断/循环三种结构、放报告里要隐去网格线。

所有长度单位为 cm，字号单位为 pt。
"""

import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SONG = "宋体"
HEI = "黑体"
KAI = "楷体"
TNR = "Times New Roman"
MONO = "Consolas"

BODY_PT = 12.0        # 小四
BODY_LINE_PT = 18.0   # 固定行距（与模板一致）
TABLE_PT = 10.5       # 五号
CAPTION_PT = 10.5
CODE_PT = 9.0

TEXT_WIDTH_CM = 21.0 - 3.17 * 2      # ≈14.66 cm，插图最大宽度
HEADER_TEXT = "华南师范大学计算机学院《编译原理》课程设计"


# ---------------------------------------------------------------------------
# 低层工具
# ---------------------------------------------------------------------------
def _set_rfonts(run, ascii_font, east_font):
    run.font.name = ascii_font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_font)
    rfonts.set(qn("w:cs"), ascii_font)


def style_run(run, east=SONG, ascii_=TNR, size=BODY_PT, bold=False,
              italic=False, color=None, underline=False):
    _set_rfonts(run, ascii_, east)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    return run


def _shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _para_border_bottom(paragraph, sz=6, color="808080"):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def _add_field(paragraph, instr, placeholder="", east=SONG, ascii_=TNR, size=BODY_PT):
    """插入一个域（PAGE / TOC 等）。"""
    run = paragraph.add_run()
    style_run(run, east=east, ascii_=ascii_, size=size)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText")
    instr_el.set(qn("xml:space"), "preserve")
    instr_el.text = instr
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text_el = OxmlElement("w:t")
    text_el.text = placeholder
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for el in (fld_begin, instr_el, fld_sep, text_el, fld_end):
        run._r.append(el)
    return run


# ---------------------------------------------------------------------------
# 文档骨架
# ---------------------------------------------------------------------------
def new_document(update_fields=True):
    doc = Document()

    # 页面：A4，与模板一致
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(3.17)
    sec.right_margin = Cm(3.17)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.header_distance = Cm(1.5)
    sec.footer_distance = Cm(1.75)

    # 正文样式：宋体 / 小四 / 固定行距 18 磅
    normal = doc.styles["Normal"]
    normal.font.name = TNR
    normal.font.size = Pt(BODY_PT)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), SONG)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(BODY_LINE_PT)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)

    # 标题样式：黑体，并显式去掉默认模板的蓝色
    specs = {
        "Heading 1": (16.0, 18.0, 12.0),   # 三号
        "Heading 2": (14.0, 12.0, 6.0),    # 四号
        "Heading 3": (12.0, 9.0, 6.0),     # 小四
        "Heading 4": (12.0, 6.0, 3.0),
    }
    for name, (size, before, after) in specs.items():
        st = doc.styles[name]
        st.font.name = TNR
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:ascii"), TNR)
        rfonts.set(qn("w:hAnsi"), TNR)
        rfonts.set(qn("w:eastAsia"), HEI)
        spf = st.paragraph_format
        spf.space_before = Pt(before)
        spf.space_after = Pt(after)
        spf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        spf.line_spacing = Pt(max(size + 6.0, BODY_LINE_PT))
        spf.keep_with_next = True
        spf.first_line_indent = Pt(0)

    add_header_footer(doc)
    if update_fields:
        settings = doc.settings.element
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        settings.append(upd)
    return doc


def add_header_footer(doc):
    sec = doc.sections[0]
    header = sec.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(hp.add_run(HEADER_TEXT), east=SONG, size=9.0)
    _para_border_bottom(hp, sz=4)

    footer = sec.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(fp.add_run("第 "), east=SONG, size=9.0)
    _add_field(fp, " PAGE ", "1", size=9.0)
    style_run(fp.add_run(" 页"), east=SONG, size=9.0)


def page_break(doc):
    from docx.enum.text import WD_BREAK
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
# 段落
# ---------------------------------------------------------------------------
_BOLD_RE = re.compile(r"\*\*(.+?)\*\*")


def _add_marked_text(paragraph, text, east, size, bold=False):
    """支持 **加粗** 标记的文本写入。"""
    pos = 0
    for m in _BOLD_RE.finditer(text):
        if m.start() > pos:
            style_run(paragraph.add_run(text[pos:m.start()]), east=east, size=size, bold=bold)
        style_run(paragraph.add_run(m.group(1)), east=east, size=size, bold=True)
        pos = m.end()
    if pos < len(text):
        style_run(paragraph.add_run(text[pos:]), east=east, size=size, bold=bold)


def p(doc, text="", indent=True, align="justify", size=BODY_PT, east=SONG,
      bold=False, space_before=0.0, space_after=0.0, keep_with_next=False):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(BODY_LINE_PT)
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.first_line_indent = Pt(size * 2) if indent else Pt(0)
    pf.keep_with_next = keep_with_next
    pf.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "right": WD_ALIGN_PARAGRAPH.RIGHT,
    }[align]
    _add_marked_text(par, text, east=east, size=size, bold=bold)
    return par


def h1(doc, text):
    par = doc.add_heading(level=1)
    style_run(par.add_run(text), east=HEI, size=16.0, bold=True)
    return par


def h2(doc, text):
    par = doc.add_heading(level=2)
    style_run(par.add_run(text), east=HEI, size=14.0, bold=True)
    return par


def h3(doc, text):
    par = doc.add_heading(level=3)
    style_run(par.add_run(text), east=HEI, size=12.0, bold=True)
    return par


def h4(doc, text):
    par = doc.add_heading(level=4)
    style_run(par.add_run(text), east=SONG, size=12.0, bold=True)
    return par


def item(doc, text, size=BODY_PT, left_pt=24.0, hang_pt=0.0):
    """带悬挂缩进的条目（用于参考文献、用例清单）。"""
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(BODY_LINE_PT)
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.left_indent = Pt(left_pt)
    pf.first_line_indent = Pt(-hang_pt if hang_pt else 0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _add_marked_text(par, text, east=SONG, size=size)
    return par


def caption(doc, text):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(15.0)
    pf.space_before = Pt(3)
    pf.space_after = Pt(6)
    pf.first_line_indent = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.keep_with_next = False
    style_run(par.add_run(text), east=SONG, size=CAPTION_PT)
    return par


def table(doc, headers, rows, widths=None, font_size=TABLE_PT, align_center=True,
          header_fill="D9E2F3", first_col_center=False):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    hdr = tbl.rows[0]
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        cell.text = ""
        par = cell.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.first_line_indent = Pt(0)
        par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        style_run(par.add_run(str(text)), east=SONG, size=font_size, bold=True)
        _shade_cell(cell, header_fill)
        cell.vertical_alignment = 1  # center

    for row in rows:
        cells = tbl.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = ""
            par = cell.paragraphs[0]
            centered = align_center and (i > 0 or first_col_center)
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
            par.paragraph_format.first_line_indent = Pt(0)
            par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            _add_marked_text(par, str(value), east=SONG, size=font_size)
            cell.vertical_alignment = 1

    if widths:
        for i, w in enumerate(widths):
            for row in tbl.rows:
                row.cells[i].width = Cm(w)
    return tbl


def figure(doc, path, width_cm, cap_text):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(6)
    pf.space_after = Pt(3)
    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf.keep_with_next = True
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    par.add_run().add_picture(path, width=Cm(min(width_cm, TEXT_WIDTH_CM)))
    caption(doc, cap_text)


def code_block(doc, code, size=CODE_PT, caption_text=None):
    """等宽字体代码块，浅灰底纹。"""
    if caption_text:
        caption(doc, caption_text)
    for line in code.rstrip("\n").split("\n"):
        par = doc.add_paragraph()
        pf = par.paragraph_format
        pf.first_line_indent = Pt(0)
        pf.left_indent = Cm(0.4)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(size + 3.0)
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_pr = par._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "F2F2F2")
        p_pr.append(shd)
        style_run(par.add_run(line if line else " "), east=SONG, ascii_=MONO, size=size)


def toc(doc):
    par = doc.add_paragraph()
    pf = par.paragraph_format
    pf.first_line_indent = Pt(0)
    pf.space_before = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _add_field(par, r' TOC \o "1-3" \h \z \u ',
               "（目录将在 Word 中自动更新；若未更新，请全选后按 F9）")


# ---------------------------------------------------------------------------
# 封面
# ---------------------------------------------------------------------------
def cover(doc, kind, subtitle, info_rows, date_text):
    """kind：文档类型（如“课程设计实验报告”）；subtitle：题目。"""
    for _ in range(2):
        p(doc, "", indent=False)

    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.first_line_indent = Pt(0)
    par.paragraph_format.space_after = Pt(18)
    style_run(par.add_run("华南师范大学计算机学院"), east=HEI, size=22.0, bold=True)

    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.first_line_indent = Pt(0)
    par.paragraph_format.space_after = Pt(6)
    style_run(par.add_run("《编译原理》课程设计"), east=HEI, size=26.0, bold=True)

    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.first_line_indent = Pt(0)
    par.paragraph_format.space_after = Pt(30)
    style_run(par.add_run(kind), east=HEI, size=26.0, bold=True)

    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.first_line_indent = Pt(0)
    par.paragraph_format.space_after = Pt(4)
    style_run(par.add_run(subtitle), east=KAI, size=15.0, bold=True)

    for _ in range(2):
        p(doc, "", indent=False)

    tbl = doc.add_table(rows=0, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for label, value in info_rows:
        cells = tbl.add_row().cells
        cells[0].width = Cm(3.6)
        cells[1].width = Cm(7.4)
        par = cells[0].paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        par.paragraph_format.first_line_indent = Pt(0)
        style_run(par.add_run(label), east=HEI, size=14.0)
        par = cells[1].paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.LEFT
        par.paragraph_format.first_line_indent = Pt(0)
        # 值域下方加下划线
        _para_border_bottom(par, sz=6)
        style_run(par.add_run(" " + value), east=SONG, size=14.0)

    for _ in range(3):
        p(doc, "", indent=False)
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.first_line_indent = Pt(0)
    style_run(par.add_run(date_text), east=HEI, size=14.0)


def info_table(doc, rows, label_width=3.2, value_width=8.0, font_size=TABLE_PT):
    tbl = doc.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for label, value in rows:
        cells = tbl.add_row().cells
        cells[0].width = Cm(label_width)
        cells[1].width = Cm(value_width)
        par = cells[0].paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        par.paragraph_format.first_line_indent = Pt(0)
        par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        style_run(par.add_run(label), east=SONG, size=font_size, bold=True)
        _shade_cell(cells[0], "F2F2F2")
        par = cells[1].paragraphs[0]
        par.paragraph_format.first_line_indent = Pt(0)
        par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        style_run(par.add_run(value), east=SONG, size=font_size)
    return tbl


# ---------------------------------------------------------------------------
# 图 / 表自动编号（按章）
# ---------------------------------------------------------------------------
class Num:
    """编号器：图 3-1、表 3-2 ……；chapter 为空时输出“图 1 / 表 1”的全局序号。
    编号由生成脚本统一分配，避免手写错号。"""

    def __init__(self, chapter=""):
        self.chapter = chapter
        self._fig = 0
        self._tab = 0

    def fig_no(self):
        self._fig += 1
        return f"图 {self.chapter}-{self._fig}" if self.chapter else f"图 {self._fig}"

    def tab_no(self):
        self._tab += 1
        return f"表 {self.chapter}-{self._tab}" if self.chapter else f"表 {self._tab}"


def figure_n(doc, num, path, title, width_cm=TEXT_WIDTH_CM):
    no = num.fig_no()
    figure(doc, path, width_cm, f"{no}  {title}")
    return no


def table_n(doc, num, title, headers, rows, widths=None, font_size=TABLE_PT,
            align_center=True, first_col_center=False, note=None):
    no = num.tab_no()
    caption(doc, f"{no}  {title}")
    tbl = table(doc, headers, rows, widths=widths, font_size=font_size,
                align_center=align_center, first_col_center=first_col_center)
    if note:
        par = doc.add_paragraph()
        par.paragraph_format.first_line_indent = Pt(0)
        par.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        par.paragraph_format.space_after = Pt(6)
        style_run(par.add_run(note), east=SONG, size=9.0)
    return no
